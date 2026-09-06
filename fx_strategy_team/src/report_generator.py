"""
report_generator.py - レポート生成モジュール
バックテスト結果をExcel形式でまとめる。
"""

import os
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import LineChart, Reference
from openpyxl.utils import get_column_letter

OUTPUT_DIR = os.path.join(
    os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
    "output", "reports"
)

HEADER_FONT = Font(bold=True, size=11, color="FFFFFF")
HEADER_FILL = PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")
TITLE_FONT = Font(bold=True, size=14, color="2F5496")
DATA_FONT = Font(size=10)
GOOD_FILL = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
BAD_FILL = PatternFill(start_color="F8CECC", end_color="F8CECC", fill_type="solid")
THIN_BORDER = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"), bottom=Side(style="thin")
)


def generate_report(pair: str, metrics: dict, trades: list, equity_curve: list,
                    is_sample_data: bool = False) -> str:
    os.makedirs(OUTPUT_DIR, exist_ok=True)
    wb = Workbook()

    _create_summary_sheet(wb, pair, metrics, is_sample_data)
    _create_trades_sheet(wb, trades)
    _create_equity_curve_sheet(wb, equity_curve)

    if "Sheet" in wb.sheetnames:
        del wb["Sheet"]

    filename = f"backtest_{pair}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    filepath = os.path.join(OUTPUT_DIR, filename)
    wb.save(filepath)
    return filepath


def _create_summary_sheet(wb, pair, metrics, is_sample_data):
    ws = wb.active
    ws.title = "サマリー"

    ws.merge_cells("A1:C1")
    ws["A1"] = f"バックテスト結果サマリー - {pair}"
    ws["A1"].font = TITLE_FONT

    if is_sample_data:
        ws.merge_cells("A2:C2")
        ws["A2"] = "※架空のサンプルデータによる動作確認用。実運用の成績を示すものではありません。"
        ws["A2"].font = Font(size=9, color="CC3300", bold=True)

    rows = [
        ("総トレード数", metrics.get("total_trades")),
        ("勝率(%)", _fmt(metrics.get("win_rate"))),
        ("プロフィットファクター", _fmt(metrics.get("profit_factor"))),
        ("期待値(1トレード平均損益)", _fmt(metrics.get("expectancy"))),
        ("最大ドローダウン(%)", _fmt(metrics.get("max_drawdown_pct"))),
        ("最終資金", _fmt(metrics.get("final_equity"))),
        ("総リターン(%)", _fmt(metrics.get("total_return_pct"))),
    ]

    row = 4
    for label, value in rows:
        ws.cell(row=row, column=1, value=label).font = Font(bold=True, size=10)
        cell = ws.cell(row=row, column=2, value=value)
        cell.font = DATA_FONT
        for c in (1, 2):
            ws.cell(row=row, column=c).border = THIN_BORDER
        row += 1

    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 18


def _create_trades_sheet(wb, trades):
    ws = wb.create_sheet(title="トレード一覧")
    headers = ["エントリー日", "決済日", "方向", "エントリー価格",
               "決済価格", "数量", "損益", "R倍数"]
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.border = THIN_BORDER

    row = 2
    for t in trades:
        values = [
            t["entry_date"], t["exit_date"], t["side"],
            round(t["entry_price"], 4), round(t["exit_price"], 4),
            round(t["units"], 2), round(t["pnl"], 2), round(t["r_multiple"], 2),
        ]
        for col, val in enumerate(values, 1):
            cell = ws.cell(row=row, column=col, value=val)
            cell.font = DATA_FONT
            cell.border = THIN_BORDER
            if col == 7:
                cell.fill = GOOD_FILL if t["pnl"] > 0 else BAD_FILL
        row += 1

    for i, w in enumerate([12, 12, 8, 14, 14, 10, 12, 10], 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _create_equity_curve_sheet(wb, equity_curve):
    ws = wb.create_sheet(title="資金推移")
    ws.cell(row=1, column=1, value="日付").font = HEADER_FONT
    ws.cell(row=1, column=2, value="資金").font = HEADER_FONT

    row = 2
    for point in equity_curve:
        ws.cell(row=row, column=1, value=point["date"])
        ws.cell(row=row, column=2, value=round(point["equity"], 2))
        row += 1

    if len(equity_curve) >= 2:
        chart = LineChart()
        chart.title = "資金推移"
        chart.width = 20
        chart.height = 10
        data = Reference(ws, min_col=2, min_row=1, max_row=row - 1)
        cats = Reference(ws, min_col=1, min_row=2, max_row=row - 1)
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(cats)
        ws.add_chart(chart, "D2")

    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 14


def print_summary(pair: str, metrics: dict, is_sample_data: bool = False):
    print("\n" + "=" * 60)
    print(f"  バックテスト結果サマリー - {pair}")
    if is_sample_data:
        print("  ※架空のサンプルデータによる動作確認用（実運用成績ではない）")
    print("=" * 60)
    print(f"  総トレード数:         {metrics.get('total_trades')}")
    print(f"  勝率:                 {_fmt(metrics.get('win_rate'))}%")
    print(f"  プロフィットファクター: {_fmt(metrics.get('profit_factor'))}")
    print(f"  期待値(1トレード平均): {_fmt(metrics.get('expectancy'))}")
    print(f"  最大ドローダウン:      {_fmt(metrics.get('max_drawdown_pct'))}%")
    print(f"  最終資金:              {_fmt(metrics.get('final_equity'))}")
    print(f"  総リターン:            {_fmt(metrics.get('total_return_pct'))}%")
    print("=" * 60 + "\n")


def _fmt(val):
    if val is None:
        return "−"
    if isinstance(val, (int, float)):
        return round(val, 2)
    return val
